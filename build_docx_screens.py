import docx
from docx.shared import Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

with open("Documentacion_Sistema.md", "r", encoding="utf-8") as f:
    doc_md = f.read()

doc = docx.Document("proyecto_gsell.docx")

settings = doc.settings.element
update_fields = OxmlElement('w:updateFields')
update_fields.set(qn('w:val'), 'true')
settings.append(update_fields)

table_style = doc.tables[-1].style if doc.tables else "Table Grid"

doc.add_page_break()

doc.add_heading('6. Diseño', level=1)
doc.add_heading('6.1. Diagrama Entidad-Relación', level=2)
doc.add_paragraph('En desarrollo.')

doc.add_heading('6.2. Frontend / Prototipo', level=2)
doc.add_paragraph('A continuación, se presentan las capturas de pantalla de las principales interfaces del sistema desarrollado en funcionamiento:')

doc.add_heading('Pantalla de Inicio de Sesión', level=3)
doc.add_picture('screen_login.png', width=Inches(6.0))

doc.add_heading('Dashboard Principal', level=3)
doc.add_picture('screen_dashboard.png', width=Inches(6.0))

doc.add_heading('Vista Pública: Reserva Online', level=3)
doc.add_picture('screen_reserva.png', width=Inches(6.0))

doc.add_heading('Gestión de Agenda de Turnos', level=3)
doc.add_picture('screen_turnos.png', width=Inches(6.0))

doc.add_heading('Gestión de Clientes', level=3)
doc.add_picture('screen_clientes.png', width=Inches(6.0))

doc.add_heading('6.3. Tabla de trazabilidad', level=2)
doc.add_paragraph('A continuación, se detalla el estado actual de los requerimientos funcionales definidos en el punto 5.7, con su correspondiente nivel de implementación en el sistema.')

data = [
    ["F-01", "Iniciar sesion", "app.py", "login()", "Implementado", ""],
    ["F-02", "Cerrar sesion (logout)", "app.py", "logout()", "Implementado", ""],
    ["F-03", "Crear usuarios", "app.py", "usuario_nuevo()", "Implementado", ""],
    ["F-04", "Editar y eliminar usuarios", "-", "-", "No implementado", "Fuera de alcance"],
    ["F-05", "Diferenciar roles", "app.py", "admin_required()", "Implementado", ""],
    ["F-06", "Registrar nuevos clientes", "app.py", "cliente_nuevo()", "Implementado", ""],
    ["F-07", "Editar datos de cliente", "app.py", "cliente_editar()", "Implementado", ""],
    ["F-08", "Eliminar clientes", "app.py", "cliente_eliminar()", "Implementado", ""],
    ["F-09", "Buscar clientes", "app.py", "clientes()", "Implementado", ""],
    ["F-10", "Guardar telefono", "app.py", "cliente_nuevo()", "Implementado", ""],
    ["F-11", "Tipo de vehiculo", "app.py", "vehiculo_nuevo()", "Implementado", ""],
    ["F-12", "Registrar vehiculos con datos", "app.py", "vehiculo_nuevo()", "Implementado", ""],
    ["F-13", "Asociar vehiculos a cliente", "app.py", "vehiculo_nuevo()", "Implementado", ""],
    ["F-14", "Editar vehiculo", "-", "-", "No implementado", "Prioridad menor"],
    ["F-15", "Historial", "app.py", "cliente_detalle()", "Implementado", ""],
    ["F-16", "Crear servicios", "app.py", "servicio_nuevo()", "Implementado", ""],
    ["F-17", "Editar servicios", "app.py", "servicio_editar()", "Implementado", ""],
    ["F-18", "Eliminar servicios", "app.py", "servicio_eliminar()", "Implementado", ""],
    ["F-19", "Asignar precios", "app.py", "servicio_nuevo()", "Implementado", ""],
    ["F-20", "Listar servicios", "app.py", "servicios()", "Implementado", ""],
    ["F-21", "Crear turnos", "app.py", "turno_nuevo()", "Implementado", ""],
    ["F-22", "Modificar turnos", "-", "-", "No implementado", "Cancelar y crear nuevo"],
    ["F-23", "Cancelar turnos", "app.py", "turno_id_estado()", "Implementado", ""],
    ["F-24", "Ver agenda", "app.py", "turnos()", "Implementado", ""],
    ["F-25", "Evitar superposicion", "app.py", "turno_nuevo()", "Implementado", ""],
    ["F-26", "Asignar cliente/vehiculo", "app.py", "turno_nuevo()", "Implementado", ""],
    ["F-27", "Asignar servicio", "app.py", "turno_nuevo()", "Implementado", ""],
    ["F-28", "Registrar pago", "app.py", "pago_nuevo()", "Implementado", ""],
    ["F-29", "Metodo de pago", "app.py", "pago_nuevo()", "Implementado", ""],
    ["F-30", "Historial pagos", "app.py", "pagos()", "Implementado", ""],
    ["F-31", "Estado pago", "app.py", "pago_nuevo()", "Implementado", ""],
    ["F-32", "Ingresos diarios", "app.py", "dashboard()", "Implementado", ""],
    ["F-33", "Ingresos mensuales", "app.py", "reportes()", "Parcial", "Solo ultimos 7 dias"],
    ["F-34", "Autos atendidos", "app.py", "dashboard()", "Implementado", ""],
    ["F-35", "Servicios mas usados", "app.py", "dashboard()", "Implementado", ""],
    ["F-36", "Resumen", "app.py", "dashboard()", "Implementado", ""],
    ["F-37", "Recordatorios", "-", "-", "No implementado", "Falta integracion externa"],
    ["F-38", "Avisos turno", "-", "-", "No implementado", "Falta integracion externa"],
    ["F-39", "Cancelaciones", "app.py", "dashboard()", "Parcial", "Solo visual en dashboard"],
    ["F-40", "Generar ticket", "app.py", "comprobante_pago()", "Implementado", ""],
    ["F-41", "Detalle servicio", "app.py", "comprobante_pago()", "Implementado", ""],
    ["F-42", "Imprimir", "app.py", "comprobante_pago()", "Parcial", "Usa impresion web"],
    ["F-43", "Buscar turnos", "app.py", "turnos()", "Parcial", "Solo filtros basicos"],
    ["F-44", "Filtrar cliente", "app.py", "clientes()", "Implementado", ""],
    ["F-45", "Filtrar servicio", "-", "-", "No implementado", "No es necesario aun"],
    ["F-46", "Buscar patente", "app.py", "vehiculos()", "Implementado", ""],
    ["F-47", "Horarios", "app.py", "configuracion()", "Implementado", ""],
    ["F-48", "Precios", "app.py", "configuracion()", "Implementado", ""],
    ["F-49", "Tipos servicios", "app.py", "configuracion()", "Implementado", ""],
    ["F-50", "Duracion servicios", "app.py", "configuracion()", "Implementado", ""]
]

table = doc.add_table(rows=1, cols=6)
table.style = table_style
hdr_cells = table.rows[0].cells
headers = ["Req", "Descripcion", "Archivo", "Funcion", "Estado", "Observacion"]
for i, h in enumerate(headers):
    hdr_cells[i].text = h

for row in data:
    row_cells = table.add_row().cells
    for i, item in enumerate(row):
        row_cells[i].text = item

doc.add_page_break()

doc.add_heading('7. Conclusión sobre el Proyecto', level=1)
doc.add_paragraph('En desarrollo.')

doc.add_heading('8. Bibliografía', level=1)
doc.add_paragraph('Sin iniciar.')

doc.add_heading('9. Glosario de palabras técnicas', level=1)
doc.add_paragraph('Sin iniciar.')

doc.add_page_break()
doc.add_heading('10. Documentación Funcional del Sistema', level=1)

for line in doc_md.split('\n'):
    if line.startswith('### '):
        doc.add_heading(line[4:], level=3)
    elif line.startswith('## '):
        doc.add_heading(line[3:], level=2)
    elif line.startswith('# '):
        doc.add_heading(line[2:], level=1)
    elif line.startswith('* '):
        doc.add_paragraph('• ' + line[2:])
    elif line.strip() != "":
        doc.add_paragraph(line)

doc.save("proyecto_gsell_final_v3.docx")
print("Document generated successfully with screenshots.")
